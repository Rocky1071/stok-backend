"""Excel export: one .xlsx with three sheets — Ürünler, Stok Hareketleri, Özel Alanlar."""

from io import BytesIO

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from lib.db import db
from lib.units import UNIT_LABELS, format_qty

router = APIRouter(prefix="/export", tags=["export"])

TYPE_LABELS = {"giris": "Giriş", "satis": "Satış", "kullanim": "Kullanım", "transfer": "Transfer"}


def _sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) if i - 1 < len(r) else 0 for r in rows] or [10]) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(width, 42)
    ws.freeze_panes = "A2"


@router.get("/excel")
async def export_excel():
    products = await db.products.find().sort("created_at", -1).to_list(1000)
    movements = await db.movements.find().sort([("date", -1), ("created_at", -1)]).to_list(1000)
    fields = await db.custom_fields.find().sort("created_at", 1).to_list(1000)

    wb = Workbook()

    ws = wb.active
    ws.title = "Ürünler"
    headers = ["Ürün Adı", "Marka", "Model", "Lot No", "Birim", "Alış Fiyatı", "Satış Fiyatı", "Stok", "Stok (Yazı)", "Kritik Stok"]
    headers += [f["label"] for f in fields]
    headers += ["Resim Sayısı", "Kapak Resmi", "Notlar", "Oluşturulma", "Son Güncelleme"]
    rows = []
    for p in products:
        unit = p.get("unit", "adet")
        qty = float(p.get("quantity", 0) or 0)
        row = [
            p.get("name", ""),
            p.get("brand", ""),
            p.get("model", ""),
            p.get("lot_number", ""),
            UNIT_LABELS.get(unit, unit),
            p.get("purchase_price", 0),
            p.get("sale_price", 0),
            qty,
            format_qty(qty, unit),
            float(p.get("min_stock", 0) or 0),
        ]
        row += [str(p.get("attributes", {}).get(f["key"], "")) for f in fields]
        images = p.get("images") or ([p["image"]] if p.get("image") else [])
        created = p.get("created_at")
        updated = p.get("updated_at")
        row += [
            len(images),
            images[0] if images else "",
            p.get("notes", ""),
            created.strftime("%d.%m.%Y %H:%M") if created else "",
            updated.strftime("%d.%m.%Y %H:%M") if updated else "",
        ]
        rows.append(row)
    _sheet(ws, headers, rows)

    ws2 = wb.create_sheet("Stok Hareketleri")
    mv_rows = [
        [
            m.get("date", ""),
            m.get("product_name", ""),
            TYPE_LABELS.get(m.get("type", ""), m.get("type", "")),
            float(m.get("quantity", 0) or 0),
            format_qty(float(m.get("quantity", 0) or 0), m.get("unit", "adet")),
            m.get("unit_price") if m.get("unit_price") is not None else "",
            m.get("destination", ""),
            m.get("note", ""),
            m.get("created_by", ""),
        ]
        for m in movements
    ]
    _sheet(
        ws2,
        ["Tarih", "Ürün", "İşlem Türü", "Miktar", "Miktar (Yazı)", "Birim Fiyat", "Hedef/Açıklama", "Not", "Kullanıcı"],
        mv_rows,
    )

    ws3 = wb.create_sheet("Özel Alanlar")
    _sheet(ws3, ["Etiket", "Anahtar", "Tip"], [[f["label"], f["key"], f["type"]] for f in fields])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="stoktakip-rapor.xlsx"'},
    )
